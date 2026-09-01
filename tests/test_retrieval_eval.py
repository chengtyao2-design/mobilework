from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

APP_ROOT = Path(__file__).resolve().parents[1]
EVAL_PATH = APP_ROOT / "benchmarks" / "retrieval_eval" / "run.py"

spec = importlib.util.spec_from_file_location("retrieval_eval", EVAL_PATH)
assert spec is not None and spec.loader is not None
retrieval_eval = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = retrieval_eval
spec.loader.exec_module(retrieval_eval)
CORPUS_ROOT = retrieval_eval.DEFAULT_CORPUS


def test_questions_manifest_covers_real_evidence_pages():
    if not CORPUS_ROOT.is_dir():
        import pytest

        pytest.skip(
            "specialized retrieval-evaluation corpus is not installed; "
            "set MOBILEWORK_EVAL_CORPUS to enable this integration check"
        )
    questions = retrieval_eval.load_questions(retrieval_eval.DEFAULT_QUESTIONS)

    assert len(questions) == 20
    assert len({question["id"] for question in questions}) == 20
    retrieval_eval.verify_expected_paths(questions, CORPUS_ROOT)
    assert {"low", "medium", "high"} <= {question["difficulty"] for question in questions}


def test_corpus_snapshot_renames_only_the_duplicate_copy(tmp_path: Path):
    source = tmp_path / "source"
    (source / "wiki" / "concepts").mkdir(parents=True)
    (source / "wiki" / "sources").mkdir(parents=True)
    (source / "raw" / "sources").mkdir(parents=True)
    (source / "wiki" / "concepts" / "海洋渔业产业集群.md").write_text("# 概念", encoding="utf-8")
    original_source = source / retrieval_eval.DUPLICATE_SOURCE_PAGE
    original_source.write_text("# 来源", encoding="utf-8")
    link_page = source / "wiki" / "concepts" / "链接.md"
    link_page.write_text(retrieval_eval.DUPLICATE_SOURCE_LINK, encoding="utf-8")

    snapshot = tmp_path / "snapshot"
    manifest = retrieval_eval.prepare_corpus_snapshot(source, snapshot)

    assert original_source.is_file()
    assert not (snapshot / retrieval_eval.DUPLICATE_SOURCE_PAGE).exists()
    assert (snapshot / retrieval_eval.SNAPSHOT_SOURCE_PAGE).is_file()
    assert retrieval_eval.SNAPSHOT_SOURCE_LINK in (snapshot / "wiki" / "concepts" / "链接.md").read_text(encoding="utf-8")
    assert manifest["source_corpus"] == str(source)


def test_parse_events_extracts_answer_citations_and_mcp_calls():
    events = "\n".join(
        [
            json.dumps(
                {
                    "type": "tool_call",
                    "toolName": "mobile-retrieval_retrieve",
                    "arguments": {"query": "一般计税方法"},
                },
                ensure_ascii=False,
            ),
            json.dumps(
                {
                    "type": "message",
                    "role": "assistant",
                    "content": "应纳税额为销项税额减进项税额。见 wiki/concepts/一般计税方法.md。",
                },
                ensure_ascii=False,
            ),
        ]
    )

    answer, citations, tool_calls, note = retrieval_eval.parse_events(events)

    assert answer.startswith("应纳税额")
    assert citations == ["wiki/concepts/一般计税方法.md"]
    assert tool_calls == ['mobile-retrieval_retrieve {"query": "一般计税方法"}']
    assert not note


def test_parse_events_marks_non_json_output():
    answer, citations, tool_calls, note = retrieval_eval.parse_events("not JSON")

    assert answer == ""
    assert citations == []
    assert tool_calls == []
    assert note == "no JSON events in OpenCode output"


def test_mcp_tool_errors_detects_failed_retrieval_call():
    event = json.dumps(
        {
            "part": {
                "type": "tool",
                "tool": "invalid",
                "state": {
                    "status": "completed",
                    "input": {
                        "tool": "mobile-retrieval_retrieve",
                        "error": "unavailable tool",
                    },
                },
            }
        },
        ensure_ascii=False,
    )

    assert retrieval_eval.mcp_tool_errors(event) == ["unavailable tool"]


def test_workbook_has_expected_sheets_and_sixty_answer_rows(tmp_path: Path):
    questions = retrieval_eval.load_questions(retrieval_eval.DEFAULT_QUESTIONS)
    attempts = []
    for question in questions:
        for tier in retrieval_eval.TIERS:
            attempts.append(
                retrieval_eval.Attempt(
                    question_id=question["id"],
                    tier=tier,
                    attempt=1,
                    status="ok",
                    started_at="2026-09-01T00:00:00+00:00",
                    finished_at="2026-09-01T00:00:01+00:00",
                    duration_seconds=1.0,
                    answer="测试答案，见 wiki/concepts/一般计税方法.md。",
                    citations=["wiki/concepts/一般计税方法.md"],
                    tool_calls=["mobile-retrieval_retrieve"],
                    exit_code=0,
                    artifact_path="run/Q01-low-attempt1.jsonl",
                )
            )

    destination = retrieval_eval.write_workbook(
        tmp_path / "evaluation.xlsx",
        {
            "run_id": "20260901T000000Z",
            "source_corpus": str(CORPUS_ROOT),
            "snapshot_corpus": str(tmp_path / "corpus"),
            "index_meta": {"index_built_at": "2026-09-01T00:00:00+0000", "rows": 1, "model": "test"},
            "opencode_version": "test",
            "model": "test-model",
            "started_at": "2026-09-01T00:00:00+00:00",
            "finished_at": "2026-09-01T00:00:01+00:00",
        },
        questions,
        attempts,
        attempts,
    )
    workbook = load_workbook(destination, data_only=False)

    assert workbook.sheetnames == ["Run Summary", "Questions", "Answers", "Failures"]
    assert workbook["Questions"].max_row == 21
    assert workbook["Answers"].max_row == 61
    assert workbook["Run Summary"]["B12"].value == "=COUNTA(Answers!A:A)-1"
