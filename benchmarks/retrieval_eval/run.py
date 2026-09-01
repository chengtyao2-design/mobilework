from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import signal
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wiki_retrieval import embedding

APP_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CORPUS = Path(
    os.environ.get("MOBILEWORK_EVAL_CORPUS", APP_ROOT.parent / "demo-data")
)
DEFAULT_QUESTIONS = Path(__file__).with_name("questions.v1.json")
DEFAULT_ARTIFACTS = Path(__file__).with_name("artifacts")
TIERS = ("low", "medium", "high")
RETRYABLE_STATUSES = {"timeout", "process_error", "malformed_events", "no_answer"}
MAX_EXCEL_TEXT = 32_767
CITATION_PATTERN = re.compile(r"(?:wiki|raw/sources)/[^\s\]\[()<>，。；：、,;:]+?\.md")
DUPLICATE_SOURCE_PAGE = Path("wiki/sources/海洋渔业产业集群.md")
SNAPSHOT_SOURCE_PAGE = Path("wiki/sources/海洋渔业产业集群（来源）.md")
DUPLICATE_SOURCE_LINK = "[[sources/海洋渔业产业集群]]"
SNAPSHOT_SOURCE_LINK = "[[sources/海洋渔业产业集群（来源）]]"


@dataclass
class Attempt:
    question_id: str
    tier: str
    attempt: int
    status: str
    started_at: str
    finished_at: str
    duration_seconds: float
    answer: str
    citations: list[str]
    tool_calls: list[str]
    exit_code: int | None
    artifact_path: str
    failure_detail: str = ""


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def utc_stamp(value: datetime) -> str:
    return value.isoformat(timespec="seconds")


def run_id_for(value: datetime) -> str:
    return value.strftime("%Y%m%dT%H%M%SZ")


def redact_text(value: str) -> str:
    redacted = value
    for name, secret in os.environ.items():
        if not secret or len(secret) < 6:
            continue
        upper = name.upper()
        if any(marker in upper for marker in ("KEY", "TOKEN", "SECRET", "PASSWORD")):
            redacted = redacted.replace(secret, "[redacted]")
    return redacted


def as_text(value: Any) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return "\n".join(part for item in value if (part := as_text(item)))
    if isinstance(value, dict):
        for key in ("text", "content", "output", "value"):
            if key in value:
                text = as_text(value[key])
                if text:
                    return text
    return ""


def walk(value: Any) -> Iterable[dict[str, Any]]:
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk(child)


def parse_events(raw: str) -> tuple[str, list[str], list[str], str]:
    events: list[dict[str, Any]] = []
    invalid_lines = 0
    for line in raw.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        try:
            event = json.loads(stripped)
        except json.JSONDecodeError:
            invalid_lines += 1
            continue
        if isinstance(event, dict):
            events.append(event)

    if not events:
        return "", [], [], "no JSON events in OpenCode output"

    full_answers: list[str] = []
    text_parts: list[str] = []
    tools: list[str] = []
    for event in events:
        for node in walk(event):
            role = str(node.get("role") or node.get("author") or "").lower()
            kind = str(node.get("type") or node.get("event") or "").lower()
            if role == "assistant":
                for key in ("content", "text", "output", "message"):
                    if key in node and (text := as_text(node[key])):
                        full_answers.append(text)
            elif kind in {"text", "text_delta", "message.delta", "assistant_text"}:
                for key in ("text", "content", "delta"):
                    if key in node and (text := as_text(node[key])):
                        text_parts.append(text)

            name = ""
            for key in ("toolName", "tool_name", "name", "tool"):
                candidate = node.get(key)
                if isinstance(candidate, dict):
                    candidate = candidate.get("name")
                if isinstance(candidate, str) and "mobile-retrieval" in candidate:
                    name = candidate
                    break
            if name:
                arguments = node.get("arguments", node.get("input", node.get("params", {})))
                detail = json.dumps(arguments, ensure_ascii=False, sort_keys=True, default=str)
                rendered = f"{name} {detail}" if detail not in ("{}", "null") else name
                if rendered not in tools:
                    tools.append(rendered)

    answer = full_answers[-1] if full_answers else "\n".join(text_parts)
    answer = redact_text(answer)
    citations = sorted(set(CITATION_PATTERN.findall(answer)))
    parse_note = ""
    if invalid_lines:
        parse_note = f"ignored {invalid_lines} non-JSON output line(s)"
    return answer, citations, tools, parse_note


def mcp_tool_errors(raw: str) -> list[str]:
    errors: list[str] = []
    for line in raw.splitlines():
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            continue
        part = event.get("part") if isinstance(event, dict) else None
        if not isinstance(part, dict) or part.get("type") != "tool":
            continue
        state = part.get("state")
        if not isinstance(state, dict):
            continue
        input_value = state.get("input")
        tool_name = part.get("tool")
        if isinstance(input_value, dict):
            tool_name = input_value.get("tool", tool_name)
        is_retrieval = isinstance(tool_name, str) and "mobile-retrieval" in tool_name
        error = ""
        if isinstance(input_value, dict):
            error = as_text(input_value.get("error"))
        if not error and state.get("status") in {"error", "failed"}:
            error = as_text(state.get("output")) or f"tool state: {state.get('status')}"
        if is_retrieval and error:
            errors.append(redact_text(error)[:1_000])
    return list(dict.fromkeys(errors))


def write_artifacts(run_dir: Path, slug: str, stdout: str, stderr: str) -> tuple[Path, Path]:
    event_path = run_dir / f"{slug}.jsonl"
    stderr_path = run_dir / f"{slug}.stderr.txt"
    event_path.write_text(redact_text(stdout), encoding="utf-8")
    stderr_path.write_text(redact_text(stderr), encoding="utf-8")
    return event_path, stderr_path


def evaluation_prompt(question: str, tier: str) -> str:
    return (
        f"请使用 wiki-ask-{tier} skill 回答以下问题。只允许调用 mobile-retrieval MCP，"
        "不要读取本地文件、执行 Shell 或引用外部知识。请以中文作答，并在每项关键结论后"
        "引用检索结果中的相对 wiki 路径；如证据不足，请明确说明缺口。\n\n"
        f"问题：{question}"
    )


def terminate_process(process: subprocess.Popen[str]) -> tuple[str, str]:
    try:
        os.killpg(process.pid, signal.SIGTERM)
    except ProcessLookupError:
        pass
    try:
        return process.communicate(timeout=10)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        return process.communicate()


def invoke_opencode(
    question: dict[str, Any],
    tier: str,
    attempt_number: int,
    corpus: Path,
    run_dir: Path,
    run_id: str,
    timeout: int,
    executable: str,
    model: str | None,
) -> Attempt:
    started = utc_now()
    slug = f"{question['id']}-{tier}-attempt{attempt_number}"
    command = [
        executable,
        "run",
        "--format",
        "json",
        "--dir",
        str(APP_ROOT),
        "--title",
        f"retrieval-eval-{run_id}-{question['id']}-{tier}",
    ]
    if model:
        command.extend(["--model", model])
    command.append(evaluation_prompt(question["question"], tier))
    environment = {**os.environ, "WIKI_RETRIEVAL_PROJECT": str(corpus)}
    stdout = ""
    stderr = ""
    exit_code: int | None = None
    status = "ok"
    failure_detail = ""
    try:
        process = subprocess.Popen(
            command,
            cwd=APP_ROOT,
            env=environment,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            start_new_session=True,
        )
        try:
            stdout, stderr = process.communicate(timeout=timeout)
        except subprocess.TimeoutExpired as error:
            partial_stdout, partial_stderr = terminate_process(process)
            stdout = partial_stdout or as_text(error.stdout)
            stderr = partial_stderr or as_text(error.stderr)
            status = "timeout"
            failure_detail = f"exceeded timeout of {timeout} seconds"
        exit_code = process.returncode
    except OSError as error:
        status = "process_error"
        failure_detail = f"could not start OpenCode: {error}"

    event_path, _ = write_artifacts(run_dir, slug, stdout, stderr)
    answer, citations, tool_calls, parse_note = parse_events(stdout)
    tool_errors = mcp_tool_errors(stdout)
    if status == "ok" and tool_errors:
        status = "retrieval_error"
        failure_detail = "\n".join(tool_errors)
    elif status == "ok" and exit_code != 0:
        status = "process_error"
        failure_detail = redact_text(stderr).strip()[:1_000] or f"OpenCode exited with {exit_code}"
    elif status == "ok" and not answer:
        status = "malformed_events" if parse_note else "no_answer"
        failure_detail = parse_note or "OpenCode completed without an assistant answer"
    elif status == "ok" and parse_note:
        failure_detail = parse_note

    finished = utc_now()
    return Attempt(
        question_id=question["id"],
        tier=tier,
        attempt=attempt_number,
        status=status,
        started_at=utc_stamp(started),
        finished_at=utc_stamp(finished),
        duration_seconds=(finished - started).total_seconds(),
        answer=answer,
        citations=citations,
        tool_calls=tool_calls,
        exit_code=exit_code,
        artifact_path=str(event_path.relative_to(run_dir.parent)),
        failure_detail=redact_text(failure_detail),
    )


def load_questions(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("schema_version") != 1:
        raise ValueError("questions manifest must use schema_version 1")
    questions = payload.get("questions")
    if not isinstance(questions, list) or len(questions) != 20:
        raise ValueError("questions manifest must contain exactly 20 questions")
    seen_ids: set[str] = set()
    required = {"id", "question", "test_point", "difficulty", "expected_paths", "evidence_note"}
    for question in questions:
        if not isinstance(question, dict) or required - question.keys():
            raise ValueError("each question must include all required fields")
        if question["id"] in seen_ids:
            raise ValueError(f"duplicate question id: {question['id']}")
        if not isinstance(question["expected_paths"], list) or not question["expected_paths"]:
            raise ValueError(f"question {question['id']} must include expected_paths")
        seen_ids.add(question["id"])
    return questions


def verify_expected_paths(questions: list[dict[str, Any]], corpus: Path) -> None:
    missing = [
        f"{question['id']}: {path}"
        for question in questions
        for path in question["expected_paths"]
        if not (corpus / path).is_file()
    ]
    if missing:
        raise FileNotFoundError("missing expected evidence pages:\n" + "\n".join(missing))


def prepare_corpus_snapshot(source: Path, destination: Path) -> dict[str, str]:
    shutil.copytree(source / "wiki", destination / "wiki")
    shutil.copytree(source / "raw" / "sources", destination / "raw" / "sources")

    duplicate_page = destination / DUPLICATE_SOURCE_PAGE
    renamed_page = destination / SNAPSHOT_SOURCE_PAGE
    if duplicate_page.is_file():
        duplicate_page.rename(renamed_page)
        for page in (destination / "wiki").rglob("*.md"):
            content = page.read_text(encoding="utf-8")
            if DUPLICATE_SOURCE_LINK in content:
                page.write_text(
                    content.replace(DUPLICATE_SOURCE_LINK, SNAPSHOT_SOURCE_LINK),
                    encoding="utf-8",
                )

    manifest = {
        "source_corpus": str(source),
        "snapshot_corpus": str(destination),
        "renamed_page": str(DUPLICATE_SOURCE_PAGE),
        "snapshot_page": str(SNAPSHOT_SOURCE_PAGE),
    }
    (destination.parent / "corpus-snapshot.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return manifest


def run_index_command(command: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(command, cwd=APP_ROOT, text=True, capture_output=True, check=False)


def ensure_index(corpus: Path, skip_index: bool) -> dict[str, Any]:
    if not skip_index:
        build = run_index_command(
            [sys.executable, "-m", "wiki_retrieval.index", "--project", str(corpus), "--batch", "16"]
        )
        if build.returncode != 0:
            raise RuntimeError(redact_text(build.stderr).strip() or "vector index build failed")

    status = run_index_command(
        [sys.executable, "-m", "wiki_retrieval.index", "--project", str(corpus), "--status"]
    )
    if status.returncode != 0:
        raise RuntimeError(redact_text(status.stderr).strip() or "vector index status check failed")
    try:
        freshness = json.loads(status.stderr.strip())
    except json.JSONDecodeError as error:
        raise RuntimeError(f"cannot parse index status: {status.stderr.strip()}") from error

    meta_path = corpus / ".lancedb" / "index_meta.json"
    if not meta_path.is_file():
        raise RuntimeError(f"index metadata missing: {meta_path}")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    if meta.get("chunking_strategy") != "whole_page_v1":
        raise RuntimeError("unexpected chunking strategy")
    if not isinstance(meta.get("rows"), int) or meta["rows"] <= 0:
        raise RuntimeError("index contains no rows")
    if freshness.get("stale_index") is not False:
        raise RuntimeError(f"index is stale: {freshness.get('reason', 'unknown reason')}")
    return {"meta": meta, "freshness": freshness}


def cell_text(value: Any, artifact_path: str = "") -> str:
    text = str(value or "")
    if len(text) <= MAX_EXCEL_TEXT:
        return text
    suffix = f"\n\n[内容超过 Excel 单元格限制，全文见 {artifact_path}]"
    return text[: MAX_EXCEL_TEXT - len(suffix)] + suffix


def style_sheet(sheet, widths: list[int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_workbook(
    destination: Path,
    run_metadata: dict[str, Any],
    questions: list[dict[str, Any]],
    canonical_attempts: list[Attempt],
    all_attempts: list[Attempt],
) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Run Summary"
    summary.append(["Field", "Value"])
    summary_rows = [
        ("Run ID", run_metadata["run_id"]),
        ("Source corpus", run_metadata["source_corpus"]),
        ("Evaluated snapshot", run_metadata["snapshot_corpus"]),
        ("Index built at", run_metadata["index_meta"].get("index_built_at", "")),
        ("Index rows", run_metadata["index_meta"].get("rows", "")),
        ("Embedding model", run_metadata["index_meta"].get("model", "")),
        ("OpenCode version", run_metadata["opencode_version"]),
        ("Requested model", run_metadata["model"] or "(default)"),
        ("Started (UTC)", run_metadata["started_at"]),
        ("Finished (UTC)", run_metadata["finished_at"]),
        ("Canonical answers", "=COUNTA(Answers!A:A)-1"),
        ("Successful answers", '=COUNTIF(Answers!C:C,"ok")'),
        ("Failed answers", "=B12-B13"),
    ]
    for row in summary_rows:
        summary.append(row)
    style_sheet(summary, [28, 110])

    questions_sheet = workbook.create_sheet("Questions")
    questions_sheet.append(
        ["Question ID", "Question", "Test point", "Difficulty", "Expected evidence paths", "Evidence note"]
    )
    for question in questions:
        questions_sheet.append(
            [
                question["id"],
                question["question"],
                question["test_point"],
                question["difficulty"],
                "\n".join(question["expected_paths"]),
                question["evidence_note"],
            ]
        )
    style_sheet(questions_sheet, [14, 48, 24, 12, 58, 68])

    answers_sheet = workbook.create_sheet("Answers")
    answers_sheet.append(
        [
            "Question ID",
            "Tier",
            "Status",
            "Attempt",
            "Started (UTC)",
            "Finished (UTC)",
            "Duration (s)",
            "Answer",
            "Citations",
            "MCP tool calls",
            "Exit code",
            "Artifact path",
            "Note",
        ]
    )
    for attempt in canonical_attempts:
        answers_sheet.append(
            [
                attempt.question_id,
                attempt.tier,
                attempt.status,
                attempt.attempt,
                attempt.started_at,
                attempt.finished_at,
                attempt.duration_seconds,
                cell_text(attempt.answer, attempt.artifact_path),
                "\n".join(attempt.citations),
                "\n".join(attempt.tool_calls),
                attempt.exit_code,
                attempt.artifact_path,
                attempt.failure_detail,
            ]
        )
    style_sheet(answers_sheet, [14, 12, 18, 10, 24, 24, 14, 90, 52, 65, 12, 50, 45])
    for cell in answers_sheet["G"][1:]:
        cell.number_format = "0.000"

    failures_sheet = workbook.create_sheet("Failures")
    failures_sheet.append(
        ["Question ID", "Tier", "Attempt", "Status", "Failure detail", "Exit code", "Artifact path"]
    )
    for attempt in all_attempts:
        if attempt.status != "ok":
            failures_sheet.append(
                [
                    attempt.question_id,
                    attempt.tier,
                    attempt.attempt,
                    attempt.status,
                    attempt.failure_detail,
                    attempt.exit_code,
                    attempt.artifact_path,
                ]
            )
    style_sheet(failures_sheet, [14, 12, 10, 18, 100, 12, 60])

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(destination)
    return destination


def opencode_version(executable: str) -> str:
    try:
        result = subprocess.run([executable, "--version"], text=True, capture_output=True, check=False)
    except OSError as error:
        return f"unavailable: {error}"
    return redact_text((result.stdout or result.stderr).strip())


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run the mobilework retrieval evaluation")
    parser.add_argument("--corpus", type=Path, default=DEFAULT_CORPUS)
    parser.add_argument("--questions", type=Path, default=DEFAULT_QUESTIONS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_ARTIFACTS)
    parser.add_argument("--excel-path", type=Path, default=None)
    parser.add_argument("--timeout", type=int, default=900)
    parser.add_argument("--skip-index", action="store_true")
    parser.add_argument("--question-id", action="append", default=[])
    parser.add_argument("--opencode", default="opencode")
    parser.add_argument("--model", default=None)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    source_corpus = args.corpus.expanduser().resolve()
    questions_path = args.questions.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    if args.timeout <= 0:
        raise ValueError("--timeout must be positive")
    if not (source_corpus / "wiki").is_dir() or not (source_corpus / "raw" / "sources").is_dir():
        raise FileNotFoundError(f"corpus must contain wiki/ and raw/sources/: {source_corpus}")

    all_questions = load_questions(questions_path)
    selected_ids = set(args.question_id)
    unknown_ids = selected_ids - {question["id"] for question in all_questions}
    if unknown_ids:
        raise ValueError(f"unknown question ids: {', '.join(sorted(unknown_ids))}")
    questions = [question for question in all_questions if not selected_ids or question["id"] in selected_ids]
    verify_expected_paths(questions, source_corpus)
    embedding.load_dotenv(APP_ROOT)

    started = utc_now()
    run_id = run_id_for(started)
    run_dir = output_dir / run_id
    run_dir.mkdir(parents=True, exist_ok=False)
    corpus = run_dir / "corpus"
    snapshot_manifest = prepare_corpus_snapshot(source_corpus, corpus)
    verify_expected_paths(questions, corpus)
    index_state = ensure_index(corpus, args.skip_index)

    canonical_attempts: list[Attempt] = []
    all_attempts: list[Attempt] = []

    for question in questions:
        for tier in TIERS:
            latest: Attempt | None = None
            for attempt_number in (1, 2):
                latest = invoke_opencode(
                    question,
                    tier,
                    attempt_number,
                    corpus,
                    run_dir,
                    run_id,
                    args.timeout,
                    args.opencode,
                    args.model,
                )
                all_attempts.append(latest)
                if latest.status == "ok" or latest.status not in RETRYABLE_STATUSES:
                    break
            if latest is None:
                raise RuntimeError("evaluation attempt was not created")
            canonical_attempts.append(latest)
            print(
                f"{latest.question_id} {latest.tier} {latest.status} "
                f"{latest.duration_seconds:.1f}s",
                flush=True,
            )

    finished = utc_now()
    destination = (
        args.excel_path.expanduser().resolve()
        if args.excel_path
        else Path.home() / "Desktop" / f"mobilework-retrieval-eval-{run_id}.xlsx"
    )
    workbook = write_workbook(
        destination,
        {
            "run_id": run_id,
            "source_corpus": snapshot_manifest["source_corpus"],
            "snapshot_corpus": snapshot_manifest["snapshot_corpus"],
            "index_meta": index_state["meta"],
            "opencode_version": opencode_version(args.opencode),
            "model": args.model,
            "started_at": utc_stamp(started),
            "finished_at": utc_stamp(finished),
        },
        questions,
        canonical_attempts,
        all_attempts,
    )
    print(workbook)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
